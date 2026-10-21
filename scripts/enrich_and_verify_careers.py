#!/usr/bin/env python3
"""
Sachin Satish Pai - 2026 Careers Tracker Enrichment & Verification Script
========================================================================
- Concurrently validates and resolves career portal redirects across 116+ companies.
- Classifies ATS systems (Workday, Greenhouse, Lever, Eightfold AI, Phenom, etc.).
- Formats Careers Sites.xlsx with dropdowns, styling, and search dashboard.
- Generates recurring automation search URLs pre-filtered for DevOps/Platform roles.
"""

import os
import sys
import json
import time
import shutil
import urllib.parse
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed
import urllib3
import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# Disable insecure HTTPS warnings for older/misconfigured certs
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

BASE_DIR = Path(__file__).resolve().parent.parent
EXCEL_PATH = BASE_DIR / "Careers Sites.xlsx"
BACKUP_PATH = BASE_DIR / "Careers Sites_backup.xlsx"
CACHE_PATH = BASE_DIR / "scripts" / "url_cache.json"

USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"

# 2026 Top Tier Platform & DevOps Employers in Bengaluru
ADDITIONAL_COMPANIES = [
    {
        "company": "Snowflake",
        "category": "Tier 1 Data Cloud",
        "url": "https://careers.snowflake.com/us/en/search-results?keywords=devops%20bengaluru",
        "target_roles": "Senior / Staff Cloud Platform Engineer",
    },
    {
        "company": "Stripe",
        "category": "Tier 1 Fintech & Infra",
        "url": "https://stripe.com/jobs/search?teams=infrastructure&location=bengaluru",
        "target_roles": "Staff Infrastructure / Platform Engineer",
    },
    {
        "company": "Rippling",
        "category": "Tier 1 High-Growth SaaS",
        "url": "https://www.rippling.com/careers?department=Engineering&location=Bengaluru%2C%20India",
        "target_roles": "Senior / Staff DevOps / Platform Engineer",
    },
    {
        "company": "Rubrik",
        "category": "Tier 1 Cloud Data Management",
        "url": "https://www.rubrik.com/company/careers",
        "target_roles": "Senior / Staff Platform & SRE Engineer",
    },
    {
        "company": "Intuit",
        "category": "Enterprise Cloud & FinTech",
        "url": "https://jobs.intuit.com/search-jobs/devops/Bengaluru",
        "target_roles": "Senior Software Engineer - Platform / DevOps",
    },
    {
        "company": "Walmart Global Tech",
        "category": "Enterprise Scale Tech",
        "url": "https://careers.walmart.com/results?q=DevOps&location=Bengaluru%2C%20KA%2C%20India",
        "target_roles": "Staff Infrastructure / DevOps Engineer",
    },
    {
        "company": "LinkedIn",
        "category": "Tier 1 Scale & Platform",
        "url": "https://careers.linkedin.com/search?keywords=DevOps%20Bangalore",
        "target_roles": "Site Reliability Engineer / Platform Engineer",
    },
    {
        "company": "Target Tech",
        "category": "Enterprise Retail Platform",
        "url": "https://jobs.target.com/search-jobs/devops/Bengaluru",
        "target_roles": "Lead / Senior DevOps Engineer",
    },
    {
        "company": "Grab",
        "category": "High-Scale Platform",
        "url": "https://grab.careers/jobs/?search=DevOps&location=bengaluru",
        "target_roles": "Senior DevOps / SRE Engineer",
    },
    {
        "company": "Coinbase",
        "category": "Tier 1 Crypto / Infra",
        "url": "https://www.coinbase.com/careers/positions?location=India",
        "target_roles": "Senior / Staff Infrastructure Engineer",
    },
    {
        "company": "Harness",
        "category": "DevOps Toolmaker (Harness CD)",
        "url": "https://www.harness.io/careers#open-positions",
        "target_roles": "Senior / Staff Platform Engineer (Harness SME)",
    },
    {
        "company": "JFrog",
        "category": "DevOps Toolmaker (Artifactory)",
        "url": "https://jfrog.com/careers/locations/india/",
        "target_roles": "Senior Cloud Platform / DevOps Engineer",
    },
    {
        "company": "Kong Inc.",
        "category": "API Gateway & Ingress Toolmaker",
        "url": "https://konghq.com/careers#open-roles",
        "target_roles": "Senior Solutions / Platform Engineer (Kong SME)",
    },
    {
        "company": "Datadog",
        "category": "Observability Toolmaker",
        "url": "https://careers.datadoghq.com/search-jobs/?location=Bengaluru%2C%20India",
        "target_roles": "Senior Cloud Platform / SRE",
    },
    {
        "company": "HashiCorp / IBM",
        "category": "Infrastructure Toolmaker",
        "url": "https://www.hashicorp.com/careers",
        "target_roles": "Senior Platform / Cloud Engineer",
    },
    {
        "company": "Couchbase",
        "category": "Distributed Database Cloud",
        "url": "https://www.couchbase.com/careers/jobs/?location=Bengaluru",
        "target_roles": "Senior Cloud / DevOps Engineer",
    },
    {
        "company": "Confluent",
        "category": "Data Streaming Cloud (Kafka)",
        "url": "https://www.confluent.io/careers/jobs/?location=Bangalore%2C+India",
        "target_roles": "Senior Cloud Infrastructure Engineer",
    },
    {
        "company": "MongoDB",
        "category": "Modern Database Cloud",
        "url": "https://www.mongodb.com/company/careers/open-roles?location=Bengaluru",
        "target_roles": "Senior Cloud / SRE Engineer",
    },
    {
        "company": "CrowdStrike",
        "category": "Cybersecurity & Cloud Platform",
        "url": "https://crowdstrike.wd5.myworkdayjobs.com/crowdstrikecareers?q=DevOps",
        "target_roles": "Senior Cloud Security & Platform Engineer",
    },
    {
        "company": "Nutanix",
        "category": "Hybrid Cloud & Infra",
        "url": "https://nutanix.wd1.myworkdayjobs.com/Nutanix?q=devops&locationCountry=e0e7a275b7c040d2a84d43e54b677a28",
        "target_roles": "Senior DevOps / Cloud Automation Engineer",
    },
    {
        "company": "Akamai Technologies",
        "category": "Cloud Edge & Delivery",
        "url": "https://akamai.wd1.myworkdayjobs.com/AkamaiCareers?q=devops",
        "target_roles": "Senior Platform / Infrastructure Engineer",
    },
    {
        "company": "Cisco / ThousandEyes",
        "category": "Enterprise Network & Observability",
        "url": "https://jobs.cisco.com/jobs/SearchJobs/DevOps?21178=%5B16948%5D",
        "target_roles": "Senior Cloud / SRE Engineer",
    },
    {
        "company": "Palo Alto Networks",
        "category": "Cybersecurity & Cloud",
        "url": "https://jobs.paloaltonetworks.com/en/jobs/?search=DevOps&country=India",
        "target_roles": "Senior DevOps / Platform Engineer",
    },
    {
        "company": "Twilio",
        "category": "Cloud Communications",
        "url": "https://www.twilio.com/en-us/company/jobs?team=Engineering&location=Bengaluru%2C+India",
        "target_roles": "Senior Cloud Infrastructure Engineer",
    },
    {
        "company": "Atlassian",
        "category": "DevOps & Collaboration",
        "url": "https://www.atlassian.com/company/careers/all-jobs?location=Bengaluru",
        "target_roles": "Senior Site Reliability Engineer / Platform",
    },
    {
        "company": "Salesforce",
        "category": "Enterprise SaaS Platform",
        "url": "https://salesforce.wd1.myworkdayjobs.com/External_Career_Site?q=devops%20bangalore",
        "target_roles": "Lead Infrastructure Engineer / SMTS DevOps",
    },
    {
        "company": "Uber",
        "category": "High-Scale Mobility & Infra",
        "url": "https://www.uber.com/in/en/careers/list/?location=IND--Bengaluru&team=Engineering",
        "target_roles": "Senior Infrastructure / Platform Engineer",
    },
    {
        "company": "Swiggy",
        "category": "High-Scale Consumer Tech",
        "url": "https://careers.swiggy.com/#/careers?keyword=devops",
        "target_roles": "Lead DevOps / Platform Engineer",
    },
    {
        "company": "Razorpay",
        "category": "High-Scale FinTech",
        "url": "https://razorpay.com/jobs/?search=DevOps",
        "target_roles": "Senior Platform / SRE Engineer",
    },
    {
        "company": "Zepto",
        "category": "Quick-Commerce High Scale",
        "url": "https://jobs.lever.co/zepto?team=Engineering",
        "target_roles": "Senior / Lead DevOps Engineer",
    },
    {
        "company": "PhonePe",
        "category": "High-Scale FinTech Infra",
        "url": "https://www.phonepe.com/careers/job-openings/?search=devops",
        "target_roles": "Staff / Lead DevOps Engineer",
    },
    {
        "company": "Postman",
        "category": "API Platform Toolmaker",
        "url": "https://boards.greenhouse.io/postman",
        "target_roles": "Senior / Staff Infrastructure & Platform Engineer",
    },
    {
        "company": "Cloudflare",
        "category": "Cloud Edge & Networking Toolmaker",
        "url": "https://boards.greenhouse.io/cloudflare",
        "target_roles": "Systems Engineer / Cloud Platform (Remote India)",
    },
    {
        "company": "GitLab",
        "category": "DevOps & CI/CD Toolmaker",
        "url": "https://boards.greenhouse.io/gitlab",
        "target_roles": "Senior Infrastructure / SRE (All-Remote)",
    },
    {
        "company": "Elastic",
        "category": "Observability & Search Toolmaker",
        "url": "https://boards.greenhouse.io/elastic",
        "target_roles": "Senior Cloud SRE / Platform Engineer",
    },
    {
        "company": "Zscaler",
        "category": "Cloud Security Giant",
        "url": "https://zscaler.wd5.myworkdayjobs.com/zscaler?q=devops%20bangalore",
        "target_roles": "Senior / Staff Cloud Operations & DevOps",
    },
    {
        "company": "Okta",
        "category": "Cloud Identity & Access Platform",
        "url": "https://okta.wd1.myworkdayjobs.com/Okta_Careers?q=devops%20bangalore",
        "target_roles": "Senior Site Reliability Engineer / Cloud Platform",
    },
    {
        "company": "Snyk",
        "category": "Developer Security & Container Sec",
        "url": "https://boards.greenhouse.io/snyk",
        "target_roles": "Senior Cloud Platform / DevOps Engineer",
    },
    {
        "company": "Dynatrace",
        "category": "Observability & APM Toolmaker",
        "url": "https://jobs.smartrecruiters.com/Dynatrace?q=devops",
        "target_roles": "Senior Infrastructure / Platform Engineer",
    },
    {
        "company": "F5 / NGINX",
        "category": "API Gateway & Ingress Toolmaker",
        "url": "https://f5.wd5.myworkdayjobs.com/F5_Careers?q=devops",
        "target_roles": "Senior Cloud Platform / SRE",
    },
    {
        "company": "Arista Networks",
        "category": "Cloud Networking & High Scale Infra",
        "url": "https://arista.wd1.myworkdayjobs.com/Arista_Careers?q=devops",
        "target_roles": "Senior Cloud Platform / DevOps Engineer",
    },
    {
        "company": "BrowserStack",
        "category": "Cloud Infrastructure Toolmaker",
        "url": "https://boards.greenhouse.io/browserstack",
        "target_roles": "Senior DevOps / Infrastructure Engineer",
    },
    {
        "company": "PayPal",
        "category": "Tier 1 FinTech & Cloud",
        "url": "https://paypal.wd1.myworkdayjobs.com/paypal-careers?q=devops%20bangalore",
        "target_roles": "MTS 1 / Senior Platform & Cloud Engineer",
    },
    {
        "company": "Pinterest",
        "category": "Tier 1 High-Scale Platform",
        "url": "https://boards.greenhouse.io/pinterest",
        "target_roles": "Senior Infrastructure / Platform Engineer",
    },
    {
        "company": "ServiceTitan",
        "category": "Enterprise Cloud SaaS",
        "url": "https://boards.greenhouse.io/servicetitan",
        "target_roles": "Senior / Staff Platform & SRE Engineer",
    },
    {
        "company": "Apple",
        "category": "Tier 1 Tech Giant",
        "url": "https://jobs.apple.com/en-in/search?search=devops&location=india-INDC",
        "target_roles": "Cloud Infrastructure Engineer / DevOps",
    },
    {
        "company": "Goldman Sachs",
        "category": "Tier 1 FinTech / Investment Bank",
        "url": "https://www.goldmansachs.com/careers/search.html",
        "target_roles": "Vice President / Senior Platform Engineer",
    },
    {
        "company": "DE Shaw & Co",
        "category": "Tier 1 Quantitative Tech",
        "url": "https://www.deshawindia.com/careers",
        "target_roles": "Senior Systems / Platform Infrastructure Engineer",
    },
    {
        "company": "Tower Research Capital",
        "category": "Tier 1 Quantitative Tech",
        "url": "https://boards.greenhouse.io/towerresearchcapital",
        "target_roles": "Senior Linux Systems & Cloud Engineer",
    },
    {
        "company": "Splunk / Cisco",
        "category": "Observability & SIEM Giant",
        "url": "https://splunk.wd1.myworkdayjobs.com/Splunk_Careers?q=devops",
        "target_roles": "Senior SRE / Cloud Platform Engineer",
    },
    {
        "company": "Qualtrics",
        "category": "Experience Cloud Platform",
        "url": "https://qualtrics.wd1.myworkdayjobs.com/Qualtrics_Careers?q=devops",
        "target_roles": "Senior DevOps / Infrastructure Engineer",
    },
    {
        "company": "Zomato / Blinkit",
        "category": "High-Scale Consumer Tech",
        "url": "https://www.zomato.com/careers",
        "target_roles": "Lead DevOps / Kubernetes Infrastructure",
    },
    {
        "company": "Navi Technologies",
        "category": "Modern FinTech Platform",
        "url": "https://navi.com/careers",
        "target_roles": "Senior / Lead Platform & SRE Engineer",
    },
    {
        "company": "Pine Labs",
        "category": "Merchant FinTech Scale",
        "url": "https://www.pinelabs.com/careers",
        "target_roles": "Senior Cloud Platform / DevOps Lead",
    },
    {
        "company": "Airmeet",
        "category": "Cloud Communications & Media",
        "url": "https://boards.greenhouse.io/airmeet",
        "target_roles": "Senior DevOps / Cloud Infrastructure",
    },
    {
        "company": "Pure Storage",
        "category": "Enterprise Cloud Data & Storage",
        "url": "https://boards.greenhouse.io/purestorage",
        "target_roles": "Senior Technical Architect / Platform Lead",
    },
    {
        "company": "New Relic",
        "category": "Observability & Telemetry Toolmaker",
        "url": "https://boards.greenhouse.io/newrelic",
        "target_roles": "Senior Cloud SRE / Platform Engineer",
    },
    {
        "company": "Sumo Logic",
        "category": "Cloud Analytics & SIEM Toolmaker",
        "url": "https://boards.greenhouse.io/sumologic",
        "target_roles": "Senior DevOps / Platform Engineer",
    },
    {
        "company": "Netskope",
        "category": "Cloud Security & SASE Platform",
        "url": "https://boards.greenhouse.io/netskope",
        "target_roles": "Staff Engineer DevOps / SRE Stack Support",
    },
    {
        "company": "Fastly",
        "category": "Edge Cloud & Networking Toolmaker",
        "url": "https://boards.greenhouse.io/fastly",
        "target_roles": "Senior Platform / Infrastructure Engineer",
    },
    {
        "company": "Starburst",
        "category": "Modern Data Lakehouse Platform",
        "url": "https://boards.greenhouse.io/starburst",
        "target_roles": "Staff AI Platform Engineer / Infrastructure Lead",
    },
    {
        "company": "InMobi",
        "category": "AdTech High-Scale Cloud Platform",
        "url": "https://boards.greenhouse.io/inmobi",
        "target_roles": "SDE III - DevOps / Lead Cloud Infrastructure",
    },
    {
        "company": "Groww",
        "category": "High-Scale FinTech Platform",
        "url": "https://boards.greenhouse.io/groww",
        "target_roles": "Lead DevOps / Cloud Platform Engineer",
    },
    {
        "company": "Meesho",
        "category": "High-Scale E-Commerce Platform",
        "url": "https://jobs.lever.co/meesho",
        "target_roles": "SDE III DevOps / Engineering Manager Platform",
    },
    {
        "company": "CRED",
        "category": "Tier 1 FinTech Platform",
        "url": "https://jobs.lever.co/cred",
        "target_roles": "Lead / Staff DevOps Engineer",
    },
    {
        "company": "Zeta",
        "category": "Next-Gen Banking Cloud Platform",
        "url": "https://jobs.lever.co/zeta",
        "target_roles": "Lead Site Reliability Engineer / Infrastructure",
    }
]

KNOWN_ATS_MAP = {
    'Snowflake': ('Workday', 'https://snowflake.wd1.myworkdayjobs.com/Snowflake_Careers?q=devops', 'https://snowflake.wd1.myworkdayjobs.com/Snowflake_Careers'),
    'Walmart Global Tech': ('Workday', 'https://walmart.wd5.myworkdayjobs.com/WalmartExternal?q=devops%20bangalore', 'https://walmart.wd5.myworkdayjobs.com/WalmartExternal'),
    'Target Tech': ('Workday', 'https://target.wd5.myworkdayjobs.com/targetcareers?q=devops%20india', 'https://target.wd5.myworkdayjobs.com/targetcareers'),
    'Intuit': ('Workday', 'https://intuit.wd5.myworkdayjobs.com/Jobs_Intuit?q=devops%20bangalore', 'https://intuit.wd5.myworkdayjobs.com/Jobs_Intuit'),
    'Grab': ('Workday', 'https://grab.wd3.myworkdayjobs.com/Grab_Careers?q=devops%20bengaluru', 'https://grab.wd3.myworkdayjobs.com/Grab_Careers'),
    'Salesforce': ('Workday', 'https://salesforce.wd1.myworkdayjobs.com/External_Career_Site?q=devops%20bangalore', 'https://salesforce.wd1.myworkdayjobs.com/External_Career_Site'),
    'CrowdStrike': ('Workday', 'https://crowdstrike.wd5.myworkdayjobs.com/crowdstrikecareers?q=devops%20india', 'https://crowdstrike.wd5.myworkdayjobs.com/crowdstrikecareers'),
    'Nutanix': ('Workday', 'https://nutanix.wd1.myworkdayjobs.com/Nutanix?q=devops', 'https://nutanix.wd1.myworkdayjobs.com/Nutanix'),
    'Akamai Technologies': ('Workday', 'https://akamai.wd1.myworkdayjobs.com/AkamaiCareers?q=devops', 'https://akamai.wd1.myworkdayjobs.com/AkamaiCareers'),
    'Palo Alto Networks': ('Workday', 'https://paloaltonetworks.wd1.myworkdayjobs.com/External?q=devops%20bangalore', 'https://paloaltonetworks.wd1.myworkdayjobs.com/External'),
    'Cisco / ThousandEyes': ('Workday', 'https://cisco.wd5.myworkdayjobs.com/Cisco_Careers?q=devops', 'https://cisco.wd5.myworkdayjobs.com/Cisco_Careers'),
    'Zscaler': ('Workday', 'https://zscaler.wd5.myworkdayjobs.com/zscaler?q=devops', 'https://zscaler.wd5.myworkdayjobs.com/zscaler'),
    'Okta': ('Workday', 'https://okta.wd1.myworkdayjobs.com/Okta_Careers?q=devops', 'https://okta.wd1.myworkdayjobs.com/Okta_Careers'),
    'PayPal': ('Workday', 'https://paypal.wd1.myworkdayjobs.com/paypal-careers?q=devops', 'https://paypal.wd1.myworkdayjobs.com/paypal-careers'),
    'F5 / NGINX': ('Workday', 'https://f5.wd5.myworkdayjobs.com/F5_Careers?q=devops', 'https://f5.wd5.myworkdayjobs.com/F5_Careers'),
    'Arista Networks': ('Workday', 'https://arista.wd1.myworkdayjobs.com/Arista_Careers?q=devops', 'https://arista.wd1.myworkdayjobs.com/Arista_Careers'),
    'Splunk / Cisco': ('Workday', 'https://splunk.wd1.myworkdayjobs.com/Splunk_Careers?q=devops', 'https://splunk.wd1.myworkdayjobs.com/Splunk_Careers'),
    'Qualtrics': ('Workday', 'https://qualtrics.wd1.myworkdayjobs.com/Qualtrics_Careers?q=devops', 'https://qualtrics.wd1.myworkdayjobs.com/Qualtrics_Careers'),
    'Stripe': ('Greenhouse', 'https://boards.greenhouse.io/stripe?department_id=4009664002', 'https://boards.greenhouse.io/stripe'),
    'Rubrik': ('Greenhouse', 'https://boards.greenhouse.io/rubrik?department_id=4025178002', 'https://boards.greenhouse.io/rubrik'),
    'Harness': ('Greenhouse', 'https://boards.greenhouse.io/harness', 'https://boards.greenhouse.io/harness'),
    'JFrog': ('Greenhouse', 'https://boards.greenhouse.io/jfrog', 'https://boards.greenhouse.io/jfrog'),
    'Kong Inc.': ('Greenhouse', 'https://boards.greenhouse.io/kong', 'https://boards.greenhouse.io/kong'),
    'Datadog': ('Greenhouse', 'https://boards.greenhouse.io/datadog', 'https://boards.greenhouse.io/datadog'),
    'Postman': ('Greenhouse', 'https://boards.greenhouse.io/postman', 'https://boards.greenhouse.io/postman'),
    'Cloudflare': ('Greenhouse', 'https://boards.greenhouse.io/cloudflare', 'https://boards.greenhouse.io/cloudflare'),
    'GitLab': ('Greenhouse', 'https://boards.greenhouse.io/gitlab', 'https://boards.greenhouse.io/gitlab'),
    'Elastic': ('Greenhouse', 'https://boards.greenhouse.io/elastic', 'https://boards.greenhouse.io/elastic'),
    'Snyk': ('Greenhouse', 'https://boards.greenhouse.io/snyk', 'https://boards.greenhouse.io/snyk'),
    'Pinterest': ('Greenhouse', 'https://boards.greenhouse.io/pinterest', 'https://boards.greenhouse.io/pinterest'),
    'ServiceTitan': ('Greenhouse', 'https://boards.greenhouse.io/servicetitan', 'https://boards.greenhouse.io/servicetitan'),
    'Tower Research Capital': ('Greenhouse', 'https://boards.greenhouse.io/towerresearchcapital', 'https://boards.greenhouse.io/towerresearchcapital'),
    'BrowserStack': ('Greenhouse', 'https://boards.greenhouse.io/browserstack', 'https://boards.greenhouse.io/browserstack'),
    'Airmeet': ('Greenhouse', 'https://boards.greenhouse.io/airmeet', 'https://boards.greenhouse.io/airmeet'),
    'Couchbase': ('Greenhouse', 'https://boards.greenhouse.io/couchbase', 'https://boards.greenhouse.io/couchbase'),
    'Confluent': ('Greenhouse', 'https://boards.greenhouse.io/confluent', 'https://boards.greenhouse.io/confluent'),
    'MongoDB': ('Greenhouse', 'https://boards.greenhouse.io/mongodb', 'https://boards.greenhouse.io/mongodb'),
    'Twilio': ('Greenhouse', 'https://boards.greenhouse.io/twilio', 'https://boards.greenhouse.io/twilio'),
    'Rippling': ('Greenhouse', 'https://boards.greenhouse.io/rippling', 'https://boards.greenhouse.io/rippling'),
    'Coinbase': ('Greenhouse', 'https://boards.greenhouse.io/coinbase', 'https://boards.greenhouse.io/coinbase'),
    'Dynatrace': ('SmartRecruiters', 'https://jobs.smartrecruiters.com/Dynatrace?q=devops', 'https://jobs.smartrecruiters.com/Dynatrace'),
    'Zepto': ('Lever', 'https://jobs.lever.co/zepto?team=Engineering', 'https://jobs.lever.co/zepto'),
    'Atlassian': ('Lever', 'https://jobs.lever.co/atlassian', 'https://jobs.lever.co/atlassian'),
    'Meesho': ('Lever', 'https://jobs.lever.co/meesho', 'https://jobs.lever.co/meesho'),
    'CRED': ('Lever', 'https://jobs.lever.co/cred', 'https://jobs.lever.co/cred'),
    'Zeta': ('Lever', 'https://jobs.lever.co/zeta', 'https://jobs.lever.co/zeta'),
    'Pure Storage': ('Greenhouse', 'https://boards.greenhouse.io/purestorage', 'https://boards.greenhouse.io/purestorage'),
    'New Relic': ('Greenhouse', 'https://boards.greenhouse.io/newrelic', 'https://boards.greenhouse.io/newrelic'),
    'Sumo Logic': ('Greenhouse', 'https://boards.greenhouse.io/sumologic', 'https://boards.greenhouse.io/sumologic'),
    'Netskope': ('Greenhouse', 'https://boards.greenhouse.io/netskope', 'https://boards.greenhouse.io/netskope'),
    'Fastly': ('Greenhouse', 'https://boards.greenhouse.io/fastly', 'https://boards.greenhouse.io/fastly'),
    'Starburst': ('Greenhouse', 'https://boards.greenhouse.io/starburst', 'https://boards.greenhouse.io/starburst'),
    'InMobi': ('Greenhouse', 'https://boards.greenhouse.io/inmobi', 'https://boards.greenhouse.io/inmobi'),
    'Groww': ('Greenhouse', 'https://boards.greenhouse.io/groww', 'https://boards.greenhouse.io/groww'),
    'Athena Health': ('Phenom / Workday', 'https://careers.athenahealth.com/us/en/search-results?keywords=devops', 'https://careers.athenahealth.com')
}


def detect_ats(url, html_sample=""):
    u = (url or "").lower()
    h = (html_sample or "").lower()
    
    if any(k in u for k in ["myworkdayjobs.com", "wd1.", "wd3.", "wd5."]) or "workday" in h:
        return "Workday"
    if "greenhouse.io" in u or "greenhouse" in h or "boards.greenhouse.io" in u:
        return "Greenhouse"
    if "lever.co" in u or "jobs.lever.co" in u or "lever" in h:
        return "Lever"
    if "eightfold.ai" in u or "eightfold" in h:
        return "Eightfold AI"
    if "smartrecruiters.com" in u or "smartrecruiters" in h:
        return "SmartRecruiters"
    if "oraclecloud.com" in u or "taleo.net" in u or "taleo" in h:
        return "Oracle Cloud / Taleo"
    if "phenompeople.com" in u or "phenom" in u or "phenom" in h:
        return "Phenom"
    if "icims.com" in u or "icims" in h:
        return "iCIMS"
    if "darwinbox.in" in u or "darwinbox" in h:
        return "Darwinbox"
    if "successfactors" in u or "successfactors" in h:
        return "SAP SuccessFactors"
    if "turbohire" in u:
        return "TurboHire"
    if "ripplehire" in u:
        return "RippleHire"
    if "ashbyhq.com" in u:
        return "Ashby"
    
    return "Direct / Custom Portal"

def build_automation_url(company, final_url, ats):
    if ats == "Workday":
        sep = "&" if "?" in final_url else "?"
        return f"{final_url}{sep}q=devops"
    elif ats == "Lever":
        sep = "&" if "?" in final_url else "?"
        return f"{final_url}{sep}team=Platform"
    elif ats == "Eightfold AI":
        sep = "&" if "?" in final_url else "?"
        return f"{final_url}{sep}query=devops&location=Bengaluru"
    elif ats == "SmartRecruiters":
        sep = "&" if "?" in final_url else "?"
        return f"{final_url}{sep}q=devops"
    return final_url

def check_single_url(session, orig_url):
    cleaned = (orig_url or "").strip()
    if not cleaned:
        return {"status": "Empty URL", "final_url": "", "ats": "None", "code": 0}
    if not cleaned.startswith("http://") and not cleaned.startswith("https://"):
        cleaned = "https://" + cleaned
        
    try:
        resp = session.get(cleaned, timeout=10, allow_redirects=True, verify=False)
        final_url = resp.url
        code = resp.status_code
        html_snippet = resp.text[:4000] if resp.text else ""
        ats = detect_ats(final_url, html_snippet)
        return {"status": f"Active ({code})", "final_url": final_url, "ats": ats, "code": code}
    except Exception as e:
        err = str(e)[:30]
        return {"status": f"Error: {err}", "final_url": cleaned, "ats": detect_ats(cleaned), "code": 0}

def main():
    print("=" * 70)
    print(" Sachin Satish Pai - Careers Tracker Enrichment Pipeline")
    print("=" * 70)

    if not EXCEL_PATH.exists():
        print(f"Error: Target file not found at {EXCEL_PATH}")
        sys.exit(1)

    # Backup
    try:
        shutil.copyfile(EXCEL_PATH, BACKUP_PATH)
        print(f"[+] Backup saved: {BACKUP_PATH.name}")
    except Exception as e:
        print(f"[-] Backup note: {e}")

    wb_orig = openpyxl.load_workbook(EXCEL_PATH)
    sheet_orig = wb_orig["Target Companies 2026"] if "Target Companies 2026" in wb_orig.sheetnames else wb_orig.active

    existing_companies = {}
    for row in sheet_orig.iter_rows(values_only=True):
        if not any(row):
            continue
        cname = str(row[0]).strip() if row[0] else ""
        if cname and cname.lower() not in ("company", "company name", "none"):
            url = str(row[8] if len(row) > 8 and row[8] else row[1]).strip()
            existing_companies[cname.lower()] = {
                "company": cname,
                "url": url,
                "category": str(row[1]) if len(row) > 1 and row[1] else "Target Employer",
                "target_roles": str(row[2]) if len(row) > 2 and row[2] else "Senior / Staff DevOps Engineer",
                "status": str(row[3]) if len(row) > 3 and row[3] else "Not Applied",
                "notes": str(row[9]) if len(row) > 9 and row[9] else ""
            }

    print(f"[+] Loaded {len(existing_companies)} existing company entries.")

    # Merge additional companies
    all_companies = list(existing_companies.values())
    for add in ADDITIONAL_COMPANIES:
        k = add["company"].lower()
        match = next((item for item in all_companies if k in item["company"].lower() or item["company"].lower() in k), None)
        if match:
            match["category"] = add["category"]
            match["target_roles"] = add["target_roles"]
        else:
            add["status"] = "Not Applied"
            add["notes"] = "Top 2026 Bengaluru Platform Target"
            all_companies.append(add)

    # Cache management
    CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
    cache = {}
    if CACHE_PATH.exists():
        try:
            with open(CACHE_PATH, "r", encoding="utf-8") as f:
                cache = json.load(f)
        except Exception:
            cache = {}

    to_check = [item for item in all_companies if item["company"] not in cache]
    if to_check:
        print(f"[*] Verifying {len(to_check)} URLs across 12 concurrent workers...")
        session = requests.Session()
        session.headers.update({"User-Agent": USER_AGENT})
        with ThreadPoolExecutor(max_workers=12) as executor:
            future_to_item = {executor.submit(check_single_url, session, item["url"]): item for item in to_check}
            for future in as_completed(future_to_item):
                item = future_to_item[future]
                cname = item["company"]
                try:
                    res = future.result()
                except Exception:
                    res = {"status": "Error", "final_url": item["url"], "ats": detect_ats(item["url"]), "code": 0}
                cache[cname] = res

        with open(CACHE_PATH, "w", encoding="utf-8") as f:
            json.dump(cache, f, indent=2)
        print("[+] URL verification complete and cached.")
    else:
        print("[+] All URLs loaded from local cache.")

    # Apply known ATS overrides
    for cname, meta in KNOWN_ATS_MAP.items():
        if cname in cache:
            cache[cname]["ats"] = meta[0]
            cache[cname]["auto_url"] = meta[1]
            cache[cname]["final_url"] = meta[2]

    # Build New Workbook
    new_wb = openpyxl.Workbook()
    ws = new_wb.active
    ws.title = "Target Companies 2026"
    ws.views.sheetView[0].showGridLines = True

    headers = [
        "Company Name", "Category / Tier", "Target Role(s)", "Application Status",
        "Date Applied", "ATS / Portal Type", "Automation Search URL (DevOps/India)",
        "Careers Portal URL", "Original URL", "Referral Notes & Contacts", "Health / HTTP Status"
    ]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )

    for col in range(1, len(headers) + 1):
        cell = ws.cell(1, col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    ws.row_dimensions[1].height = 28

    def sort_rank(item):
        if "athena" in item["company"].lower(): return (0, item["company"])
        if "Toolmaker" in item.get("category", ""): return (1, item["company"])
        if "Tier 1" in item.get("category", ""): return (2, item["company"])
        return (3, item["company"])

    sorted_list = sorted(all_companies, key=sort_rank)

    fill_current = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
    fill_alt = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    bold_font = Font(name="Calibri", size=10, bold=True, color="000000")
    reg_font = Font(name="Calibri", size=10, color="000000")
    link_font = Font(name="Calibri", size=10, color="0563C1", underline="single")

    for r_idx, item in enumerate(sorted_list, start=2):
        cname = item["company"]
        cat = item.get("category", "Target Employer")
        roles = item.get("target_roles", "Senior / Staff DevOps Engineer")
        status = item.get("status", "Not Applied")
        orig_url = item.get("url", "")
        notes = item.get("notes", "")

        c_meta = cache.get(cname, {})
        ats = c_meta.get("ats", detect_ats(orig_url))
        final_url = c_meta.get("final_url", orig_url)
        auto_url = c_meta.get("auto_url", build_automation_url(cname, final_url, ats))
        h_status = c_meta.get("status", "Active (200)")

        row_vals = [cname, cat, roles, status, "", ats, auto_url, final_url, orig_url, notes, h_status]
        ws.append(row_vals)
        ws.row_dimensions[r_idx].height = 20

        is_even = (r_idx % 2 == 0)
        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(r_idx, c_idx)
            cell.border = thin_border
            cell.font = reg_font
            cell.alignment = left_align
            if "Current Employer" in status:
                cell.fill = fill_current
            elif is_even:
                cell.fill = fill_alt

            if c_idx == 1: cell.font = bold_font
            elif c_idx in (4, 5, 6, 11): cell.alignment = center_align
            elif c_idx in (7, 8, 9) and str(cell.value).startswith("http"):
                cell.hyperlink = str(cell.value)
                cell.font = link_font

    # Validation
    dv = DataValidation(type="list", formula1='"Not Applied,Wishlist,Referral Requested,Applied,Interviewing,Current Employer (SMTS),Offer,Rejected"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"D2:D{len(sorted_list) + 1}")

    widths = {1: 28, 2: 28, 3: 38, 4: 24, 5: 16, 6: 22, 7: 45, 8: 45, 9: 35, 10: 36, 11: 18}
    for c_idx, w in widths.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = w
    ws.freeze_panes = "A2"

    # Dashboard sheet
    dash = new_wb.create_sheet("Search Dashboard")
    dash.views.sheetView[0].showGridLines = True
    dash.column_dimensions["A"].width = 30
    dash.column_dimensions["B"].width = 40
    dash.append(["Sachin Satish Pai - 2026 Job Search Dashboard", ""])
    dash.merge_cells("A1:B1")
    dash.cell(1, 1).font = Font(name="Calibri", size=14, bold=True, color="1B365D")
    dash.row_dimensions[1].height = 28

    dash_rows = [
        ("Candidate", "Sachin Satish Pai"),
        ("Current Role & Company", "Senior Member of Technical Staff (SMTS) @ Athenahealth"),
        ("Target Roles", "Senior / Staff DevOps Engineer, Platform Engineer, Cloud Infrastructure Lead"),
        ("Primary Location", "Bengaluru, India (Hybrid / Remote)"),
        ("Total Companies Tracked", len(sorted_list)),
        ("Workday Portals Identified", sum(1 for c in cache.values() if "Workday" in c.get("ats", ""))),
        ("Greenhouse Portals", sum(1 for c in cache.values() if "Greenhouse" in c.get("ats", ""))),
        ("Lever Portals", sum(1 for c in cache.values() if "Lever" in c.get("ats", ""))),
        ("Eightfold AI Portals", sum(1 for c in cache.values() if "Eightfold" in c.get("ats", ""))),
        ("Other / Direct Portals", sum(1 for c in cache.values() if not any(k in c.get("ats", "") for k in ["Workday", "Greenhouse", "Lever", "Eightfold"])))
    ]

    for k, v in dash_rows:
        dash.append([k, v])
        r = dash.max_row
        dash.cell(r, 1).font = bold_font
        dash.cell(r, 1).fill = fill_alt
        dash.cell(r, 1).border = thin_border
        dash.cell(r, 2).font = reg_font
        dash.cell(r, 2).border = thin_border
        dash.row_dimensions[r].height = 20

    # Save to local staging scratch file first to avoid OneDrive mini-filter locking
    temp_stage = BASE_DIR / "scripts" / "temp_stage.xlsx"
    new_wb.save(temp_stage)
    shutil.copyfile(temp_stage, EXCEL_PATH)
    if temp_stage.exists():
        temp_stage.unlink()

    print(f"[+] Successfully deployed enriched tracker to {EXCEL_PATH.name}!")
    print(f"[+] Total Tracked: {len(sorted_list)} companies across Workday, Greenhouse, Lever, and Direct portals.")

if __name__ == "__main__":
    main()
